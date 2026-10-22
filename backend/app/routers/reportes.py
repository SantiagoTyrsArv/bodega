from datetime import UTC, datetime, timedelta, timezone
from decimal import Decimal
from io import BytesIO
from typing import Annotated

from fastapi import APIRouter, Depends, HTTPException, Query, Response, status
from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet
from sqlalchemy import func, select
from sqlalchemy.orm import Session

from app.db import get_db
from app.esquemas import ProductoSalida, Resumen
from app.modelos import Motivo, Movimiento, Producto, TipoMovimiento, Unidad
from app.routers.movimientos import filtros_movimientos, listar_movimientos
from app.routers.productos import bajo_minimo_sql
from app.seguridad import UsuarioActual

router = APIRouter(tags=["reportes"])

Db = Annotated[Session, Depends(get_db)]
# Minutos respecto a UTC de quien descarga (Colombia: -300). Excel no guarda zona horaria,
# así que las fechas se escriben en la hora local de esa persona
Desfase = Annotated[int, Query(ge=-14 * 60, le=14 * 60)]

CENTAVO = Decimal("0.01")
CANTIDAD = "#,##0.000"
DINERO = "#,##0.00"
FECHA = "yyyy-mm-dd hh:mm"
TIPO_XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

TEXTO_UNIDAD = {
    Unidad.UNIDAD: "Unidad",
    Unidad.KG: "Kg",
    Unidad.LITRO: "Litro",
    Unidad.CAJA: "Caja",
}
TEXTO_TIPO = {TipoMovimiento.ENTRADA: "Entrada", TipoMovimiento.SALIDA: "Salida"}
TEXTO_MOTIVO = {
    Motivo.COMPRA: "Compra",
    Motivo.DEVOLUCION: "Devolución",
    Motivo.USO: "Uso",
    Motivo.VENTA: "Venta",
    Motivo.PERDIDA: "Pérdida",
    Motivo.AJUSTE: "Ajuste",
    Motivo.ANULACION: "Anulación",
}

# Columnas de cada reporte: (título, ancho, formato numérico)
Columnas = list[tuple[str, int, str | None]]

COLUMNAS_STOCK: Columnas = [
    ("Código", 16, None),
    ("Producto", 32, None),
    ("Presentación", 14, None),
    ("Categoría", 16, None),
    ("Unidad", 9, None),
    ("Stock", 12, CANTIDAD),
    ("Stock mínimo", 13, CANTIDAD),
    ("Bajo mínimo", 12, None),
    ("Precio compra", 14, DINERO),
    ("Precio venta", 14, DINERO),
    ("Valor a compra", 16, DINERO),
    ("Valor a venta", 16, DINERO),
]

COLUMNAS_MOVIMIENTOS: Columnas = [
    ("Fecha", 17, FECHA),
    ("N.º", 7, None),
    ("Código", 16, None),
    ("Producto", 32, None),
    ("Presentación", 14, None),
    ("Tipo", 9, None),
    ("Motivo", 12, None),
    ("Cantidad", 12, CANTIDAD),
    ("Unidad", 9, None),
    ("Precio unitario", 15, DINERO),
    ("Valor", 15, DINERO),
    ("Stock después", 14, CANTIDAD),
    ("Usuario", 18, None),
    ("N.º pedido", 12, None),
    ("Cliente / responsable", 22, None),
    ("Comentario", 30, None),
    ("Anula al N.º", 12, None),
    ("Anulado por N.º", 15, None),
]


@router.get("/resumen")
def resumen(_: UsuarioActual, db: Db) -> Resumen:
    activos = Producto.activo.is_(True)
    total, valor_compra, valor_venta = db.execute(
        select(
            func.count(),
            func.coalesce(func.sum(Producto.stock * Producto.precio_compra), 0),
            func.coalesce(func.sum(Producto.stock * Producto.precio_venta), 0),
        ).where(activos)
    ).one()
    bajo_minimo = db.scalars(
        select(Producto).where(activos, bajo_minimo_sql()).order_by(Producto.nombre)
    )
    return Resumen(
        total_productos=total,
        valor_compra=Decimal(valor_compra).quantize(CENTAVO),
        valor_venta=Decimal(valor_venta).quantize(CENTAVO),
        bajo_minimo=[ProductoSalida.model_validate(p) for p in bajo_minimo],
        ultimos_movimientos=listar_movimientos(db, por_pagina=10).items,
    )


def _hoja(libro: Workbook, titulo: str, columnas: Columnas) -> Worksheet:
    """Primera hoja del libro con encabezados en negrita, anchos y la fila 1 fija."""
    hoja = libro.active
    hoja.title = titulo
    hoja.append([c[0] for c in columnas])
    for i, (_, ancho, _) in enumerate(columnas, start=1):
        hoja.column_dimensions[get_column_letter(i)].width = ancho
        hoja.cell(row=1, column=i).font = Font(bold=True)
    hoja.freeze_panes = "A2"
    return hoja


def _agregar_fila(hoja: Worksheet, valores: list) -> None:
    """Agrega una fila. Los textos (nombres, comentarios) los escriben los usuarios: se marcan
    como texto para que uno que empiece con "=" no se convierta en fórmula al abrir el Excel."""
    hoja.append(valores)
    for celda in hoja[hoja.max_row]:
        if isinstance(celda.value, str):
            celda.data_type = "s"


def _formatos(hoja: Worksheet, columnas: Columnas) -> None:
    for i, (_, _, formato) in enumerate(columnas, start=1):
        if formato:
            for (celda,) in hoja.iter_rows(min_row=2, min_col=i, max_col=i):
                celda.number_format = formato


def _respuesta(libro: Workbook, nombre: str) -> Response:
    salida = BytesIO()
    libro.save(salida)
    return Response(
        salida.getvalue(),
        media_type=TIPO_XLSX,
        headers={"Content-Disposition": f'attachment; filename="{nombre}"'},
    )


@router.get("/reportes/stock.xlsx")
def stock_xlsx(_: UsuarioActual, db: Db, desfase: Desfase = 0) -> Response:
    """Productos activos con stock, precios y valor; al final, la fila de totales."""
    libro = Workbook()
    hoja = _hoja(libro, "Stock", COLUMNAS_STOCK)
    productos = db.scalars(
        select(Producto).where(Producto.activo.is_(True)).order_by(Producto.nombre)
    )
    total_compra = total_venta = Decimal(0)
    for p in productos:
        total_compra += p.stock * p.precio_compra
        total_venta += p.stock * p.precio_venta
        _agregar_fila(
            hoja,
            [
                p.codigo,
                p.nombre,
                p.presentacion or "",
                p.categoria.nombre if p.categoria else "",
                TEXTO_UNIDAD[Unidad(p.unidad)],
                p.stock,
                p.stock_minimo,
                "Sí" if p.stock_minimo > 0 and p.stock <= p.stock_minimo else "No",
                p.precio_compra,
                p.precio_venta,
                (p.stock * p.precio_compra).quantize(CENTAVO),
                (p.stock * p.precio_venta).quantize(CENTAVO),
            ],
        )
    # Igual que /resumen: se redondea la suma, no cada fila
    hoja.append(
        ["Total"] + [None] * 9 + [total_compra.quantize(CENTAVO), total_venta.quantize(CENTAVO)]
    )
    for celda in hoja[hoja.max_row]:
        celda.font = Font(bold=True)
    _formatos(hoja, COLUMNAS_STOCK)

    hoy = (datetime.now(UTC) + timedelta(minutes=desfase)).date()
    return _respuesta(libro, f"stock-{hoy}.xlsx")


@router.get("/reportes/movimientos.xlsx")
def movimientos_xlsx(
    _: UsuarioActual,
    db: Db,
    desde: datetime | None = None,
    hasta: datetime | None = None,
    tipo: TipoMovimiento | None = None,
    desfase: Desfase = 0,
) -> Response:
    """Movimientos entre desde (inclusivo) y hasta (exclusivo), del más antiguo al más reciente."""
    if desde is not None and hasta is not None and desde >= hasta:
        raise HTTPException(
            status.HTTP_422_UNPROCESSABLE_CONTENT, "La fecha inicial debe ser anterior a la final"
        )
    libro = Workbook()
    hoja = _hoja(libro, "Movimientos", COLUMNAS_MOVIMIENTOS)
    zona = timezone(timedelta(minutes=desfase))
    consulta = (
        select(Movimiento)
        .where(*filtros_movimientos(desde=desde, hasta=hasta, tipo=tipo))
        .order_by(Movimiento.fecha, Movimiento.id)
    )
    for m in db.scalars(consulta):
        _agregar_fila(
            hoja,
            [
                m.fecha.astimezone(zona).replace(tzinfo=None),
                m.id,
                m.producto.codigo,
                m.producto.nombre,
                m.producto.presentacion or "",
                TEXTO_TIPO[TipoMovimiento(m.tipo)],
                TEXTO_MOTIVO[Motivo(m.motivo)],
                m.cantidad,
                TEXTO_UNIDAD[Unidad(m.producto.unidad)],
                m.precio_unitario,
                (m.cantidad * m.precio_unitario).quantize(CENTAVO),
                m.stock_despues,
                m.usuario.nombre,
                m.pedido or "",
                m.responsable or "",
                m.comentario or "",
                m.anula_a,
                m.anulado_por,
            ],
        )
    _formatos(hoja, COLUMNAS_MOVIMIENTOS)
    hoja.auto_filter.ref = hoja.dimensions

    # El nombre lleva el rango en días locales; hasta es exclusivo, por eso se resta un instante
    def dia(f: datetime) -> str:
        return f.astimezone(zona).date().isoformat()

    partes = []
    if desde:
        partes.append(f"desde-{dia(desde)}")
    if hasta:
        partes.append(f"hasta-{dia(hasta - timedelta(microseconds=1))}")
    return _respuesta(libro, f"movimientos-{'_'.join(partes) or 'todos'}.xlsx")
