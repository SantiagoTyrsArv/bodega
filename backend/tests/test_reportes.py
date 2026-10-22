from datetime import datetime
from decimal import Decimal
from io import BytesIO

import pytest
from openpyxl import load_workbook

from app.modelos import Movimiento, Rol


@pytest.fixture
def admin(crear_usuario, headers_de):
    crear_usuario("admin2@prueba.local", rol=Rol.ADMIN)
    return headers_de("admin2@prueba.local")


@pytest.fixture
def usuario(crear_usuario, headers_de):
    crear_usuario("ana@prueba.local")
    return headers_de("ana@prueba.local")


def crear(cliente, admin, codigo, compra, venta, minimo="0"):
    datos = {
        "codigo": codigo,
        "nombre": f"Producto {codigo}",
        "unidad": "KG",
        "precio_compra": compra,
        "precio_venta": venta,
        "stock_minimo": minimo,
    }
    r = cliente.post("/api/productos", json=datos, headers=admin)
    assert r.status_code == 201, r.text
    return r.json()


def mover(cliente, headers, producto_id, tipo, cantidad, **extra):
    motivo = "COMPRA" if tipo == "ENTRADA" else "USO"
    datos = {"producto_id": producto_id, "tipo": tipo, "motivo": motivo, "cantidad": cantidad}
    r = cliente.post("/api/movimientos", json={**datos, **extra}, headers=headers)
    assert r.status_code == 201, r.text
    return r.json()


def filas(respuesta):
    assert respuesta.status_code == 200, respuesta.text
    assert respuesta.headers["content-type"].startswith(
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    hoja = load_workbook(BytesIO(respuesta.content)).active
    return [list(f) for f in hoja.iter_rows(values_only=True)]


def test_stock_xlsx_con_valores_y_total(cliente, admin, usuario):
    a = crear(cliente, admin, "A1", "2500.00", "3200.00", minimo="20")
    mover(cliente, usuario, a["id"], "ENTRADA", "10.5")
    b = crear(cliente, admin, "B2", "1999.99", "2500.50")
    mover(cliente, usuario, b["id"], "ENTRADA", "3.125")
    inactivo = crear(cliente, admin, "C3", "100.00", "100.00")
    cliente.patch(f"/api/productos/{inactivo['id']}", json={"activo": False}, headers=admin)

    r = cliente.get("/api/reportes/stock.xlsx?desfase=-300", headers=usuario)
    datos = filas(r)
    assert "attachment" in r.headers["content-disposition"]
    assert 'filename="stock-' in r.headers["content-disposition"]

    assert datos[0][:5] == ["Código", "Producto", "Categoría", "Unidad", "Stock"]
    assert [f[0] for f in datos[1:]] == ["A1", "B2", "Total"]  # sin el desactivado
    a1 = datos[1]
    assert a1[4] == 10.5 and a1[6] == "Sí"
    assert a1[9] == 26250  # 10.5 × 2500
    # El total coincide con /resumen
    resumen = cliente.get("/api/resumen", headers=usuario).json()
    total = datos[-1]
    assert Decimal(str(total[9])) == Decimal(resumen["valor_compra"])
    assert Decimal(str(total[10])) == Decimal(resumen["valor_venta"])


def test_movimientos_xlsx_filtra_por_fecha_y_usa_hora_local(cliente, admin, usuario, db):
    p = crear(cliente, admin, "A1", "100.00", "150.00")
    m1 = mover(cliente, usuario, p["id"], "ENTRADA", "5", comentario="primero")
    m2 = mover(cliente, usuario, p["id"], "SALIDA", "2")
    anulacion = cliente.post(f"/api/movimientos/{m2['id']}/anular", headers=admin).json()

    # Fechas fijas para probar el filtro y la conversión a la hora de Colombia (UTC-5)
    fechas = {
        m1["id"]: "2026-10-01T04:30:00+00:00",  # 30 sep 23:30 en Colombia
        m2["id"]: "2026-10-01T15:00:00+00:00",
        anulacion["id"]: "2026-10-02T15:00:00+00:00",
    }
    for id_, fecha in fechas.items():
        db.get(Movimiento, id_).fecha = datetime.fromisoformat(fecha)
    db.commit()

    # Todo el 1 de octubre en hora de Colombia
    params = {"desde": "2026-10-01T00:00:00-05:00", "hasta": "2026-10-02T00:00:00-05:00"}
    r = cliente.get(
        "/api/reportes/movimientos.xlsx", params={**params, "desfase": -300}, headers=usuario
    )
    datos = filas(r)
    assert (
        'filename="movimientos-desde-2026-10-01_hasta-2026-10-01.xlsx"'
        in r.headers["content-disposition"]
    )
    assert [f[1] for f in datos[1:]] == [m2["id"]]  # m1 cae el 30 de sep en Colombia
    fila = datos[1]
    assert fila[0] == datetime(2026, 10, 1, 10, 0)
    assert fila[4:8] == ["Salida", "Uso", 2, "Kg"]
    assert fila[9] == 300  # 2 × 150
    assert fila[14] == anulacion["id"]  # anulado por

    # Sin filtros: los tres, del más antiguo al más reciente
    datos = filas(cliente.get("/api/reportes/movimientos.xlsx", headers=usuario))
    assert [f[1] for f in datos[1:]] == [m1["id"], m2["id"], anulacion["id"]]
    assert datos[1][0] == datetime(2026, 10, 1, 4, 30)  # sin desfase, en UTC
    assert datos[1][12] == "primero"
    assert datos[3][5] == "Anulación" and datos[3][13] == m2["id"]

    # Por tipo
    datos = filas(cliente.get("/api/reportes/movimientos.xlsx?tipo=ENTRADA", headers=usuario))
    assert [f[1] for f in datos[1:]] == [m1["id"], anulacion["id"]]


def test_movimientos_xlsx_rango_invertido(cliente, usuario):
    params = {"desde": "2026-10-02T00:00:00Z", "hasta": "2026-10-01T00:00:00Z"}
    r = cliente.get("/api/reportes/movimientos.xlsx", params=params, headers=usuario)
    assert r.status_code == 422
    assert r.json()["detail"] == "La fecha inicial debe ser anterior a la final"


def test_reportes_requieren_sesion(cliente):
    assert cliente.get("/api/reportes/stock.xlsx").status_code == 401
    assert cliente.get("/api/reportes/movimientos.xlsx").status_code == 401


def test_textos_que_parecen_formulas_se_guardan_como_texto(cliente, admin, usuario):
    peligroso = '=HYPERLINK("http://ejemplo.com","clic")'
    datos = {
        "codigo": "=1+2",
        "nombre": peligroso,
        "unidad": "UNIDAD",
        "precio_compra": "1.00",
        "precio_venta": "1.00",
    }
    p = cliente.post("/api/productos", json=datos, headers=admin).json()
    mover(cliente, usuario, p["id"], "ENTRADA", "1", comentario="=SUM(A1:A9)")

    for ruta, columnas in [
        ("/api/reportes/stock.xlsx", {1: "=1+2", 2: peligroso}),
        ("/api/reportes/movimientos.xlsx", {3: "=1+2", 4: peligroso, 13: "=SUM(A1:A9)"}),
    ]:
        r = cliente.get(ruta, headers=usuario)
        hoja = load_workbook(BytesIO(r.content)).active
        for columna, esperado in columnas.items():
            celda = hoja.cell(row=2, column=columna)
            assert celda.value == esperado
            assert celda.data_type == "s", f"{ruta} columna {columna} quedó como fórmula"
